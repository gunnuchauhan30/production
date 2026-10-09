"""
Reports (v2.2 addition).

Two things live here:
  1. GET /reports/summary   -> JSON, powers the in-app Reports dashboard.
  2. GET /reports/export    -> a single .xlsx workbook download with one
     sheet per report (OEE - Tube Line, OEE - Fabrication Line,
     Production Sheet, Press Stock Summary).

Both accept optional date_from / date_to query params (inclusive).
Everything is computed live from the same tables the rest of the app
writes to -- there is no separate "reporting" data store.
"""
import io
from collections import defaultdict
from datetime import date
from decimal import Decimal
from typing import Optional

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from ..database import get_db
from ..auth import get_current_user
from .. import models, schemas

router = APIRouter(prefix="/reports", tags=["reports"])

TUBE_LINE = "Tube Line"
FAB_LINE = "Fabrication Line"


def _avg(values):
    vals = [v for v in values if v is not None]
    if not vals:
        return None
    return round(sum(Decimal(v) for v in vals) / len(vals), 5)


def _oee_summary_rows(db: Session, line: str, date_from, date_to):
    query = db.query(models.ProductionOeeEntry).filter(models.ProductionOeeEntry.line == line)
    if date_from:
        query = query.filter(models.ProductionOeeEntry.entry_date >= date_from)
    if date_to:
        query = query.filter(models.ProductionOeeEntry.entry_date <= date_to)
    entries = query.all()

    grouped = defaultdict(list)
    for e in entries:
        grouped[(e.operator_name, e.machine_no)].append(e)

    rows = []
    for (operator_name, machine_no), items in grouped.items():
        rows.append(
            schemas.OeeLineSummaryRow(
                operator_name=operator_name,
                machine_no=machine_no,
                entries_count=len(items),
                total_ok_qty=sum((it.ok_qty or 0) for it in items),
                total_rejected_qty=sum((it.rejected_qty or 0) for it in items),
                avg_availability=_avg([it.availability for it in items]),
                avg_performance=_avg([it.performance for it in items]),
                avg_quality=_avg([it.quality for it in items]),
                avg_oee=_avg([it.oee for it in items]),
            )
        )
    rows.sort(key=lambda r: (r.operator_name or ""))
    return rows


def _production_summary_rows(db: Session, date_from, date_to):
    query = db.query(models.ProductionEntry, models.RouteCard.part_number).join(
        models.RouteCard, models.ProductionEntry.route_card_id == models.RouteCard.id
    )
    if date_from:
        query = query.filter(models.ProductionEntry.entry_date >= date_from)
    if date_to:
        query = query.filter(models.ProductionEntry.entry_date <= date_to)
    results = query.all()

    grouped = defaultdict(lambda: {"qty_produced": 0, "ok_qty": 0, "rejected_qty": 0, "rework_qty": 0})
    for entry, part_number in results:
        key = (entry.entry_date, part_number, entry.operation_name)
        grouped[key]["qty_produced"] += entry.qty_produced or 0
        grouped[key]["ok_qty"] += entry.ok_qty or 0
        grouped[key]["rejected_qty"] += entry.rejected_qty or 0
        grouped[key]["rework_qty"] += entry.rework_qty or 0

    rows = [
        schemas.ProductionSummaryRow(
            entry_date=k[0], part_number=k[1], operation_name=k[2], **v
        )
        for k, v in grouped.items()
    ]
    rows.sort(key=lambda r: (r.entry_date,))
    return rows


def _press_stock_summary_rows(db: Session, date_from, date_to):
    query = db.query(models.PressStockEntry)
    if date_from:
        query = query.filter(models.PressStockEntry.entry_date >= date_from)
    if date_to:
        query = query.filter(models.PressStockEntry.entry_date <= date_to)
    entries = query.all()

    grouped = defaultdict(list)
    for e in entries:
        grouped[(e.machine_no, e.operator_name)].append(e)

    rows = []
    for (machine_no, operator_name), items in grouped.items():
        rows.append(
            schemas.PressStockSummaryRow(
                machine_no=machine_no,
                operator_name=operator_name,
                entries_count=len(items),
                total_strokes=sum((it.strokes or 0) for it in items),
                total_calculation=sum((it.calculation or 0) for it in items),
            )
        )
    rows.sort(key=lambda r: (r.machine_no or ""))
    return rows


@router.get("/summary", response_model=schemas.ReportsSummary)
def get_reports_summary(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    return schemas.ReportsSummary(
        date_from=date_from,
        date_to=date_to,
        oee_tube_line=_oee_summary_rows(db, TUBE_LINE, date_from, date_to),
        oee_fabrication_line=_oee_summary_rows(db, FAB_LINE, date_from, date_to),
        production=_production_summary_rows(db, date_from, date_to),
        press_stock=_press_stock_summary_rows(db, date_from, date_to),
    )


# ---------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------
HEADER_FILL = PatternFill(start_color="C0182B", end_color="C0182B", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="C0182B")


def _write_sheet(ws, title, headers, rows):
    ws["A1"] = "CHANDA ENTERPRISES PVT. LTD."
    ws["A1"].font = TITLE_FONT
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=11)
    header_row = 4
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for r_idx, row in enumerate(rows, start=header_row + 1):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
    for col, h in enumerate(headers, start=1):
        width = max(14, len(str(h)) + 4)
        ws.column_dimensions[chr(64 + col) if col <= 26 else "A"].width = width


def _pct(v):
    return f"{float(v) * 100:.1f}%" if v is not None else ""


@router.get("/export")
def export_reports_excel(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "OEE - Tube Line"
    tube_rows = _oee_summary_rows(db, TUBE_LINE, date_from, date_to)
    _write_sheet(
        ws1,
        "OEE Summary — Tube Line",
        ["Operator Name", "Machine No", "Entries", "OK Qty", "Rejected Qty", "Availability", "Performance", "Quality", "OEE"],
        [
            [r.operator_name, r.machine_no, r.entries_count, float(r.total_ok_qty), float(r.total_rejected_qty),
             _pct(r.avg_availability), _pct(r.avg_performance), _pct(r.avg_quality), _pct(r.avg_oee)]
            for r in tube_rows
        ],
    )

    ws2 = wb.create_sheet("OEE - Fabrication Line")
    fab_rows = _oee_summary_rows(db, FAB_LINE, date_from, date_to)
    _write_sheet(
        ws2,
        "OEE Summary — Fabrication Line",
        ["Operator Name", "Machine No", "Entries", "OK Qty", "Rejected Qty", "Availability", "Performance", "Quality", "OEE"],
        [
            [r.operator_name, r.machine_no, r.entries_count, float(r.total_ok_qty), float(r.total_rejected_qty),
             _pct(r.avg_availability), _pct(r.avg_performance), _pct(r.avg_quality), _pct(r.avg_oee)]
            for r in fab_rows
        ],
    )

    ws3 = wb.create_sheet("Production Sheet")
    prod_rows = _production_summary_rows(db, date_from, date_to)
    _write_sheet(
        ws3,
        "Production Sheet",
        ["Date", "Part Number", "Operation", "Qty Produced", "OK Qty", "Rejected Qty", "Rework Qty"],
        [
            [r.entry_date.strftime("%d-%m-%Y"), r.part_number, r.operation_name,
             float(r.qty_produced), float(r.ok_qty), float(r.rejected_qty), float(r.rework_qty)]
            for r in prod_rows
        ],
    )

    ws4 = wb.create_sheet("Press Stock Summary")
    press_rows = _press_stock_summary_rows(db, date_from, date_to)
    _write_sheet(
        ws4,
        "Press Shop — No. of Stocks (Strokes)",
        ["Machine No", "Operator Name", "Entries", "Total Strokes", "Total Calculation"],
        [
            [r.machine_no, r.operator_name, r.entries_count, float(r.total_strokes), float(r.total_calculation)]
            for r in press_rows
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname_bits = ["Chanda_PPC_Report"]
    if date_from:
        fname_bits.append(date_from.strftime("%d%m%Y"))
    if date_to:
        fname_bits.append(date_to.strftime("%d%m%Y"))
    filename = "_".join(fname_bits) + ".xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
